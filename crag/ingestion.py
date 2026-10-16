"""Ingestion pipeline: upload -> extract -> clean -> split."""
from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}


@dataclass
class Page:
    """A unit of extracted text with its location in the source document."""
    text: str
    page: int | None = None


@dataclass
class Chunk:
    text: str
    source: str
    chunk_id: str
    page: int | None = None
    metadata: dict = field(default_factory=dict)

    def citation(self) -> str:
        return f"{self.source} (p. {self.page})" if self.page else self.source


# ----------------------------------------------------------------------------
# Extraction
# ----------------------------------------------------------------------------
def extract_pages(filename: str, data: bytes) -> list[Page]:
    """Extract raw text from a file's bytes, keeping page numbers for PDFs."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}")

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return [Page(p.extract_text() or "", page=i + 1) for i, p in enumerate(reader.pages)]

    if ext == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [Page("\n".join(parts))]

    return [Page(data.decode("utf-8", errors="ignore"))]


# ----------------------------------------------------------------------------
# Cleaning
# ----------------------------------------------------------------------------
_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")
_MULTI_SPACE = re.compile(r"[ \t\f\v]+")
_MULTI_NEWLINE = re.compile(r"\n{3,}")
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def clean_text(text: str) -> str:
    """Normalise unicode, re-join hyphenated line breaks and collapse whitespace."""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL.sub("", text)
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    text = "\n".join(_MULTI_SPACE.sub(" ", line).strip() for line in text.split("\n"))
    text = _MULTI_NEWLINE.sub("\n\n", text)
    return text.strip()


# ----------------------------------------------------------------------------
# Splitting
# ----------------------------------------------------------------------------
_SEPARATORS = ["\n\n", "\n", ". ", "? ", "! ", "; ", ", ", " "]


def _split_recursive(text: str, size: int, separators: list[str]) -> list[str]:
    if len(text) <= size:
        return [text]
    if not separators:
        return [text[i : i + size] for i in range(0, len(text), size)]

    sep, rest = separators[0], separators[1:]
    if sep not in text:
        return _split_recursive(text, size, rest)

    pieces: list[str] = []
    for part in text.split(sep):
        # keep sentence punctuation attached to the sentence it ends
        part = part + sep.strip() if sep.strip() else part
        if len(part) > size:
            pieces.extend(_split_recursive(part, size, rest))
        elif part:
            pieces.append(part)
    return pieces


def split_text(text: str, chunk_size: int = 800, chunk_overlap: int = 150) -> list[str]:
    """Recursive character splitter that prefers paragraph/sentence boundaries
    and carries `chunk_overlap` characters of context into the next chunk."""
    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap must be smaller than chunk_size")

    pieces = _split_recursive(text, chunk_size, _SEPARATORS)
    chunks: list[str] = []
    current = ""
    for piece in pieces:
        joiner = "" if not current or current.endswith((" ", "\n")) else " "
        if len(current) + len(joiner) + len(piece) <= chunk_size:
            current += joiner + piece
            continue
        if current.strip():
            chunks.append(current.strip())
        tail = current[-chunk_overlap:] if chunk_overlap else ""
        # start the overlap on a word boundary
        if " " in tail:
            tail = tail[tail.index(" ") + 1 :]
        current = (tail + " " + piece) if tail and len(tail) + len(piece) + 1 <= chunk_size else piece
    if current.strip():
        chunks.append(current.strip())
    return chunks


# ----------------------------------------------------------------------------
# End-to-end
# ----------------------------------------------------------------------------
def process_document(
    filename: str, data: bytes, chunk_size: int = 800, chunk_overlap: int = 150
) -> list[Chunk]:
    """Run the full ingestion pipeline for one uploaded file."""
    source = Path(filename).name
    doc_hash = hashlib.sha1(data).hexdigest()[:12]
    chunks: list[Chunk] = []
    for page in extract_pages(filename, data):
        cleaned = clean_text(page.text)
        if not cleaned:
            continue
        for piece in split_text(cleaned, chunk_size, chunk_overlap):
            if len(piece) < 20:  # drop fragments with no real content
                continue
            chunks.append(
                Chunk(
                    text=piece,
                    source=source,
                    page=page.page,
                    chunk_id=f"{doc_hash}-{len(chunks)}",
                    metadata={"doc_hash": doc_hash},
                )
            )
    return chunks
